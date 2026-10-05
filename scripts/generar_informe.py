"""
Genera DemandIA_Fresh_Avance1_Informe_Definitivo.docx.

Todos los números del informe se calculan aquí mismo ejecutando la simulación
del proyecto (no se escriben a mano), así informe, README y código coinciden.

Uso (desde la carpeta del proyecto):
    pip install python-docx matplotlib
    python scripts/generar_informe.py
"""

import sys
import tempfile
from dataclasses import replace
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import simulation  # noqa: E402
from config import (AJUSTE_REACTIVO, CAMBIO_NIVEL, ESTRATEGIAS, FACTOR_DIA_ALTO,  # noqa: E402
                    FACTOR_DIA_SEMANA, MAX_PRODUCCION_REL, MIN_PRODUCCION_REL,
                    PASO_CANDIDATOS, PROB_DIA_ALTO, RUIDO_DIARIO, UMBRAL_BAJAR,
                    UMBRAL_SUBIR, VENTANA_HISTORIAL, Parametros)
from metrics import calcular_metricas_finales, estrategia_ganadora  # noqa: E402

SALIDA = RAIZ / "DemandIA_Fresh_Avance1_Informe_Definitivo.docx"
INTEGRANTES = ["Santiago Valdivia", "Fabian Cristobal", "Jesus Camargo", "Said Taravay"]
COLORES = {"Base": "#8d99ae", "Reactivo": "#f4a261", "Utilidad": "#2a9d8f"}
AZUL = RGBColor(0x1D, 0x35, 0x57)

P = Parametros()


# ============================ 1) RESULTADOS REALES ============================
def calcular_resultados():
    df = simulation.ejecutar_simulacion(P)
    metricas = calcular_metricas_finales(df)
    validacion = simulation.ejecutar_multiples_corridas(P, 30)

    # Cambios en vivo: mismas condiciones salvo un parámetro
    experimentos = []
    for nombre, cambio in [("Caso original", {}),
                           ("Costo de desperdicio de S/ 2 a S/ 8", {"costo_desperdicio": 8.0}),
                           ("Penalización por venta perdida de S/ 5 a S/ 15",
                            {"penalizacion_venta_perdida": 15.0})]:
        d = simulation.ejecutar_simulacion(replace(P, **cambio))
        m = calcular_metricas_finales(d)
        prod_media = d[d["estrategia"] == "Utilidad"]["produccion"].mean()
        experimentos.append((nombre, prod_media, m))

    # Sensibilidad: qué pasa si no hay cambio de nivel de la demanda
    original = simulation.CAMBIO_NIVEL
    simulation.CAMBIO_NIVEL = 0.0
    try:
        validacion_sin_cambio = simulation.ejecutar_multiples_corridas(P, 30)
    finally:
        simulation.CAMBIO_NIVEL = original

    return df, metricas, validacion, experimentos, validacion_sin_cambio


# ================================ 2) FIGURAS ================================
def figura_flujo(ruta):
    fig, ax = plt.subplots(figsize=(8, 8.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 10.6)
    ax.axis("off")

    def caja(x, y, w, h, texto, color="#1d3557", fondo="#e9f1f7", size=11, bold=True):
        ax.add_patch(plt.Rectangle((x, y), w, h, fc=fondo, ec=color, lw=1.8, zorder=2))
        ax.text(x + w / 2, y + h / 2, texto, ha="center", va="center", fontsize=size,
                color=color, fontweight="bold" if bold else "normal", zorder=3)

    def flecha(x, y0, y1):
        ax.annotate("", xy=(x, y1), xytext=(x, y0),
                    arrowprops=dict(arrowstyle="-|>", color="#555", lw=2), zorder=1)

    caja(2, 9.4, 6, 1.0, "DATOS / DEMANDA\ngenerar_demanda()  ·  simulation.py")
    flecha(5, 9.4, 8.7)
    caja(2, 7.7, 6, 1.0, "PERCEPCIÓN\nhistorial de demanda y ventas del día anterior")
    flecha(5, 7.7, 7.0)
    caja(0.2, 5.5, 3.0, 1.5, "BASE\ndecidir_base()\nproduce una\ncantidad fija", COLORES["Base"], "#f1f3f6", 10)
    caja(3.5, 5.5, 3.0, 1.5, "REACTIVO\ndecidir_reactivo()\nreglas IF-THEN\ndel día anterior", "#c9772a", "#fdf0e2", 10)
    caja(6.8, 5.5, 3.0, 1.5, "UTILIDAD\ndecidir_utilidad()\nmaximiza la utilidad\nesperada", "#1f776d", "#e1f4f1", 10)
    for x in (1.7, 5.0, 8.3):
        flecha(x, 5.5, 4.85)
    caja(2, 3.85, 6, 1.0, "DECISIÓN DE PRODUCCIÓN\n(antes de conocer la demanda del día)")
    flecha(5, 3.85, 3.15)
    caja(2, 2.15, 6, 1.0, "SIMULACIÓN\nejecutar_simulacion()  ·  misma demanda para los 3")
    flecha(5, 2.15, 1.45)
    caja(2, 0.45, 6, 1.0, "MÉTRICAS\ncalcular_resultado_dia()  ·  calcular_metricas_finales()")
    ax.text(5, 0.1, "↓  COMPARACIÓN: gana la mayor utilidad neta acumulada", ha="center",
            fontsize=11, color="#1d3557", fontweight="bold")
    fig.savefig(ruta, dpi=170, bbox_inches="tight")
    plt.close(fig)


def figura_utilidad(df, ruta):
    fig, ax = plt.subplots(figsize=(8, 4))
    for e in ESTRATEGIAS:
        d = df[df["estrategia"] == e]
        ax.plot(d["dia"], d["utilidad_acumulada"], label=e, color=COLORES[e], lw=2.5)
    ax.set_xlabel("Día")
    ax.set_ylabel("Utilidad acumulada (S/)")
    ax.set_title("Utilidad acumulada por día")
    ax.grid(alpha=0.3)
    ax.legend()
    fig.savefig(ruta, dpi=170, bbox_inches="tight")
    plt.close(fig)


def figura_demanda_produccion(df, ruta):
    fig, ax = plt.subplots(figsize=(8, 4))
    d = df[df["estrategia"] == "Utilidad"]
    ax.plot(d["dia"], d["demanda"], label="Demanda real", color="#264653", marker="o", ms=3)
    for e in ESTRATEGIAS:
        x = df[df["estrategia"] == e]
        ax.plot(x["dia"], x["produccion"], label=f"Producción ({e})", color=COLORES[e],
                ls="--", lw=2)
    ax.set_xlabel("Día")
    ax.set_ylabel("Unidades")
    ax.set_title("Demanda real vs producción de cada estrategia")
    ax.grid(alpha=0.3)
    ax.legend(fontsize=8, ncol=2)
    fig.savefig(ruta, dpi=170, bbox_inches="tight")
    plt.close(fig)


# ============================ 3) AYUDAS PARA WORD ============================
def sombrear(celda, color_hex):
    tc_pr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color_hex)
    tc_pr.append(sh)


def tabla(doc, encabezado, filas, anchos=None, resaltar_fila=None):
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, texto in enumerate(encabezado):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(texto)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(c, "1D3557")
    for n, fila in enumerate(filas):
        cells = t.add_row().cells
        for i, texto in enumerate(fila):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(texto))
            r.font.size = Pt(9.5)
            if resaltar_fila is not None and n == resaltar_fila:
                r.bold = True
                sombrear(cells[i], "D8F3DC")
            if i > 0 and any(ch.isdigit() for ch in str(texto)) and len(str(texto)) < 16:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if anchos:
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def parrafo(doc, texto, negrita_inicio=None, italica=False):
    p = doc.add_paragraph()
    if negrita_inicio:
        p.add_run(negrita_inicio).bold = True
    r = p.add_run(texto)
    r.italic = italica
    return p


def vineta(doc, texto, negrita_inicio=None):
    p = doc.add_paragraph(style="List Bullet")
    if negrita_inicio:
        p.add_run(negrita_inicio).bold = True
    p.add_run(texto)
    return p


def numero_pagina(seccion):
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    run.font.size = Pt(9)


def formato(x):
    return f"{x:,.0f}"


# ================================ 4) DOCUMENTO ================================
def construir():
    df, m, validacion, experimentos, val_sin_cambio = calcular_resultados()
    ganador = estrategia_ganadora(m)

    # Guardas: el texto del informe describe estos resultados. Si el código o los
    # parámetros cambian y algo de esto deja de ser cierto, hay que revisar el texto.
    assert ganador == "Utilidad", "El texto asume que Utilidad gana; revisar la interpretación"
    assert m["Desperdicio total"].idxmin() == "Base"
    assert m["Desperdicio total"].idxmax() == "Reactivo"
    assert m["Nivel de servicio (%)"].idxmax() == "Reactivo"
    assert m["Nivel de servicio (%)"].idxmin() == "Base"

    u = m["Utilidad neta acumulada"]
    ventaja_base = u["Utilidad"] - u["Base"]
    ventaja_reactivo = u["Utilidad"] - u["Reactivo"]
    gan_val = validacion["Corridas ganadas (%)"]
    gan_val_sc = val_sin_cambio["Corridas ganadas (%)"]
    exp_desp, exp_pen = experimentos[1], experimentos[2]
    prod0, prod_desp, prod_pen = experimentos[0][1], exp_desp[1], exp_pen[1]
    assert prod_desp < prod0 < prod_pen, "La conducta esperada del agente no se cumple"

    tmp = Path(tempfile.mkdtemp())
    figura_flujo(tmp / "flujo.png")
    figura_utilidad(df, tmp / "utilidad.png")
    figura_demanda_produccion(df, tmp / "demanda.png")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    numero_pagina(sec)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    estilo.paragraph_format.space_after = Pt(6)
    estilo.paragraph_format.line_spacing = 1.15
    for nivel, tam in ((1, 16), (2, 13)):
        h = doc.styles[f"Heading {nivel}"]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = AZUL

    # ------------------------------- PORTADA -------------------------------
    def centrado(texto, tam, negrita=False, color=None, antes=0, despues=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(despues)
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.bold = negrita
        if color:
            r.font.color.rgb = color
        return p

    centrado("Facultad de Ingeniería e Inteligencia Artificial", 14, True, AZUL, antes=30)
    centrado("Ingeniería de Sistemas de Información", 13, False, AZUL)
    centrado("Trabajo Final — Avance 1", 13, True, antes=60, despues=40)
    centrado("DemandIA Fresh", 34, True, AZUL, despues=4)
    centrado("Agente inteligente basado en datos para reducir desperdicio y quiebres de "
             "stock en negocios de alimentos perecibles", 12, False, RGBColor(0x55, 0x55, 0x55),
             despues=50)
    centrado("Curso: Agentes Inteligentes", 12, True)
    centrado("Profesor: Dr. José Alfredo Herrera Quispe", 12)
    centrado("Semestre: 2026-2", 12, despues=30)
    centrado("Integrantes", 12, True, AZUL, despues=4)
    for nombre in INTEGRANTES:
        centrado(nombre, 12, despues=2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------ CONTENIDO ------------------------------
    secciones = ["Introducción", "Descripción del problema", "Descripción de DemandIA Fresh",
                 "Objetivo general", "Objetivos específicos", "Metodología", "Arquitectura",
                 "Estrategias comparadas", "Modo base", "Agente reactivo simple",
                 "Agente basado en utilidad", "Función de utilidad", "Métricas",
                 "Resultados reales", "Tabla comparativa", "Interpretación",
                 "Conclusiones", "Trabajo futuro", "Uso de IA", "Roles del equipo"]
    doc.add_heading("Contenido", level=1)
    for i, s in enumerate(secciones, start=1):
        p = doc.add_paragraph(f"{i}. {s}")
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    n = iter(range(1, 100))

    def h1(titulo):
        doc.add_heading(f"{next(n)}. {titulo}", level=1)

    # ----------------------------- INTRODUCCIÓN -----------------------------
    h1("Introducción")
    parrafo(doc, "Este documento presenta el Avance 1 del Trabajo Final del curso Agentes "
                 "Inteligentes. El proyecto se llama DemandIA Fresh y busca responder una "
                 "pregunta práctica: ¿puede un agente inteligente tomar mejores decisiones de "
                 "producción diaria que una regla fija, en un negocio de alimentos perecibles?")
    parrafo(doc, "Para responderla construimos un simulador en Python y Streamlit. En él, tres "
                 "estrategias deciden cuánto producir cada día y se enfrentan exactamente a la "
                 "misma demanda. Luego comparamos los resultados con una misma métrica "
                 "principal: la utilidad neta acumulada. En esta primera parte no se usa "
                 "Machine Learning; eso queda para la Parte 2.")

    # ------------------------------- PROBLEMA -------------------------------
    h1("Descripción del problema")
    parrafo(doc, "Una cafetería, una panadería o un pequeño negocio de alimentos perecibles debe "
                 "decidir cada día cuánto producir antes de conocer la demanda real. Esa "
                 "decisión tiene dos riesgos opuestos:")
    vineta(doc, " genera desperdicio, pierde dinero y gasta recursos innecesarios.",
           "Si produce demasiado,")
    vineta(doc, " pierde ventas, tiene quiebres de stock, baja su nivel de servicio y puede "
                "afectar la satisfacción del cliente.", "Si produce poco,")
    parrafo(doc, "El problema lo sufren los negocios pequeños que venden productos que no se "
                 "pueden guardar de un día para otro y que toman la decisión “a ojo”, con "
                 "información incompleta y demanda variable.")

    # ------------------------------ DESCRIPCIÓN ------------------------------
    h1("Descripción de DemandIA Fresh")
    parrafo(doc, "DemandIA Fresh es una aplicación web que simula varios días de operación de un "
                 "negocio de alimentos perecibles. En cada día, cada estrategia decide cuánto "
                 "producir sin ver la demanda de ese día; después se revela la demanda real y se "
                 "calculan ventas, desperdicio, ventas perdidas y utilidad.")
    parrafo(doc, "La interfaz permite cambiar los parámetros del negocio (precio, costos, "
                 "penalización, demanda promedio, producción fija, número de días y semilla), "
                 "ejecutar la simulación y ver al instante los indicadores, la tabla "
                 "comparativa y los gráficos.")

    h1("Objetivo general")
    parrafo(doc, "Demostrar experimentalmente, mediante una simulación reproducible, si un agente "
                 "reactivo simple y un agente basado en utilidad mejoran la decisión de "
                 "producción diaria frente a un modo base de producción fija, usando la utilidad "
                 "neta acumulada como métrica principal.")

    h1("Objetivos específicos")
    for texto in [
        "Simular una demanda diaria razonablemente realista y reproducible mediante una semilla.",
        "Implementar tres estrategias: modo base, agente reactivo simple y agente basado en utilidad.",
        "Enfrentar las tres estrategias a exactamente la misma secuencia de demanda.",
        "Calcular utilidad, ventas, producción, desperdicio, ventas perdidas y nivel de servicio.",
        "Mostrar los resultados en una aplicación Streamlit con KPIs, tabla comparativa y gráficos.",
        "Dejar el código claro para poder explicarlo y modificarlo en la sustentación.",
    ]:
        vineta(doc, texto)

    # ------------------------------ METODOLOGÍA ------------------------------
    h1("Metodología")
    parrafo(doc, "Seguimos un enfoque experimental basado en simulación. Los pasos son:")
    for i, texto in enumerate([
        f"Se genera una secuencia de {P.dias} días de demanda con una semilla fija (seed = {P.seed}).",
        "Los tres métodos reciben esa misma demanda y los mismos parámetros económicos.",
        "Cada día, cada método decide cuánto producir sin conocer la demanda de ese día.",
        "Se revela la demanda real y se calculan ventas, desperdicio, ventas perdidas y utilidad.",
        "Se acumulan los resultados y se calculan las métricas finales por estrategia.",
        "Se compara con la utilidad neta acumulada; gana la estrategia que la tenga mayor.",
        "Como validación extra, se repite el experimento con 30 semillas consecutivas.",
    ], start=1):
        parrafo(doc, texto, negrita_inicio=f"Paso {i}. ")

    parrafo(doc, "Demanda simulada (sin Machine Learning). ", None).runs[0].bold = True
    parrafo(doc, "La demanda de cada día se calcula como: demanda promedio × efecto del día de la "
                 "semana × variación aleatoria × (día de demanda alta) × (cambio de nivel). "
                 f"El efecto del día de la semana va de {min(FACTOR_DIA_SEMANA):.2f} a "
                 f"{max(FACTOR_DIA_SEMANA):.2f}; la variación aleatoria es normal con desviación "
                 f"de {RUIDO_DIARIO:.0%}; cada día tiene {PROB_DIA_ALTO:.0%} de probabilidad de "
                 f"ser de demanda alta ({FACTOR_DIA_ALTO - 1:+.0%}); y desde la mitad del periodo "
                 f"la demanda sube {CAMBIO_NIVEL:.0%} (por ejemplo, cambio de temporada). La "
                 "demanda se redondea y nunca es negativa.")

    parrafo(doc, "Parámetros usados en los resultados de este informe:")
    tabla(doc, ["Parámetro", "Valor"], [
        ["Demanda promedio", f"{P.demanda_media} unidades/día"],
        ["Precio de venta por unidad", f"S/ {P.precio_venta:.0f}"],
        ["Costo de producción por unidad", f"S/ {P.costo_produccion:.0f}"],
        ["Costo por unidad desperdiciada", f"S/ {P.costo_desperdicio:.0f}"],
        ["Penalización por venta perdida", f"S/ {P.penalizacion_venta_perdida:.0f}"],
        ["Producción fija del modo base", f"{P.produccion_base} unidades/día"],
        ["Número de días", f"{P.dias}"],
        ["Semilla aleatoria", f"{P.seed}"],
        ["Límites de producción de los agentes",
         f"{P.min_produccion} a {P.max_produccion} unidades "
         f"({MIN_PRODUCCION_REL:.0%} y {MAX_PRODUCCION_REL:.0%} de la demanda promedio)"],
    ], anchos=[7, 8.5])

    # ------------------------------ ARQUITECTURA ------------------------------
    h1("Arquitectura")
    parrafo(doc, "El sistema está dividido en pocos archivos, cada uno con una responsabilidad "
                 "clara:")
    tabla(doc, ["Archivo", "Responsabilidad", "Funciones principales"], [
        ["config.py", "Parámetros y constantes (umbrales, ventana, límites)", "Parametros"],
        ["agents.py", "Las tres estrategias de decisión",
         "decidir_base(), decidir_reactivo(), decidir_utilidad()"],
        ["simulation.py", "Demanda y ciclo diario de simulación",
         "generar_demanda(), ejecutar_simulacion()"],
        ["metrics.py", "Fórmula de utilidad y métricas",
         "calcular_resultado_dia(), calcular_metricas_finales()"],
        ["app.py", "Interfaz Streamlit", "—"],
    ], anchos=[3, 6, 6.5])
    parrafo(doc, "El flujo completo del sistema es el siguiente:")
    doc.add_picture(str(tmp / "flujo.png"), width=Cm(11.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo(doc, "Figura 1. Flujo del sistema.", italica=True).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------------------ ESTRATEGIAS ------------------------------
    h1("Estrategias comparadas")
    tabla(doc, ["Estrategia", "Tipo", "Qué percibe", "Cómo decide"], [
        ["Base", "Modo base", "Nada", "Produce siempre la misma cantidad"],
        ["Reactivo", "Agente reactivo simple", "Producción y ventas del día anterior",
         "Reglas IF-THEN sobre el porcentaje vendido"],
        ["Utilidad", "Agente basado en utilidad", f"Demandas de los últimos {VENTANA_HISTORIAL} días",
         "Evalúa cantidades posibles y elige la de mayor utilidad esperada"],
    ], anchos=[2.5, 3.5, 4.5, 5])

    h1("Modo base")
    parrafo(doc, f"El modo base no es inteligente: produce siempre {P.produccion_base} unidades "
                 "por día, sin importar lo que ocurra. Se implementa en la función "
                 "decidir_base() de agents.py. Representa a un negocio que produce “lo de "
                 "siempre” y sirve como referencia para saber si los agentes realmente aportan "
                 "una mejora.")

    h1("Agente reactivo simple")
    parrafo(doc, "El agente reactivo (función decidir_reactivo() en agents.py) usa únicamente "
                 "información inmediata: lo producido y lo vendido el día anterior. Calcula el "
                 "porcentaje vendido = ventas de ayer / producción de ayer y aplica estas reglas:")
    vineta(doc, f" si el porcentaje vendido es ≥ {UMBRAL_SUBIR:.0%}, aumenta la producción en {AJUSTE_REACTIVO:.0%}.", "IF")
    vineta(doc, f" si el porcentaje vendido es ≤ {UMBRAL_BAJAR:.0%}, reduce la producción en {AJUSTE_REACTIVO:.0%}.", "IF")
    vineta(doc, " en cualquier otro caso, mantiene la producción.", "ELSE")
    parrafo(doc, f"El primer día no hay “ayer”, así que parte de la producción base. La producción "
                 f"siempre se limita al rango de {P.min_produccion} a {P.max_produccion} "
                 "unidades. Es un agente reactivo simple porque solo asocia una condición actual "
                 "con una acción; no estima la demanda futura ni compara alternativas.")

    h1("Agente basado en utilidad")
    parrafo(doc, "El agente de utilidad (función decidir_utilidad() en agents.py) no usa Machine "
                 "Learning. Su idea es sencilla: probar varias acciones posibles y quedarse con "
                 "la que maximiza la utilidad esperada. El procedimiento es:")
    for i, texto in enumerate([
        f"Toma la demanda observada en los últimos {VENTANA_HISTORIAL} días. Esos valores son "
        f"sus {VENTANA_HISTORIAL} escenarios posibles de demanda para mañana, todos igual de "
        "probables. Si todavía no hay tantos días, completa los que faltan con la demanda "
        "promedio configurada.",
        f"Define las cantidades candidatas: de {P.min_produccion} a {P.max_produccion}, de "
        f"{PASO_CANDIDATOS} en {PASO_CANDIDATOS}.",
        "Para cada cantidad calcula la utilidad que obtendría en cada escenario y las promedia: "
        "esa es la utilidad esperada.",
        "Elige la cantidad con mayor utilidad esperada.",
    ], start=1):
        parrafo(doc, texto, negrita_inicio=f"{i}. ")
    parrafo(doc, "Como el agente compara alternativas con una medida numérica (la utilidad), se "
                 "comporta distinto según los costos: si el desperdicio es muy caro tiende a "
                 "producir menos; si perder una venta es muy caro, tiende a producir más.")

    h1("Función de utilidad")
    parrafo(doc, "La utilidad de un día se calcula en calcular_resultado_dia() de metrics.py. La "
                 "misma función la usa la simulación para medir resultados y el agente de "
                 "utilidad para estimar cada alternativa.")
    for linea in ["ventas = min(producción, demanda)",
                  "desperdicio = max(producción − demanda, 0)",
                  "ventas_perdidas = max(demanda − producción, 0)",
                  "ingresos = ventas × precio_venta",
                  "costo_producción = producción × costo_unitario",
                  "costo_desperdicio_total = desperdicio × costo_desperdicio",
                  "penalización = ventas_perdidas × penalización_venta_perdida"]:
        p = doc.add_paragraph(linea)
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(1)
        p.runs[0].font.name = "Consolas"
        p.runs[0].font.size = Pt(10)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run("Utilidad = ingresos − costo de producción − costo de desperdicio − penalización").bold = True
    parrafo(doc, f"Ejemplo con los parámetros base: si se producen 120 unidades y la demanda es "
                 f"100, se venden 100 (ingresos S/ 1,500), cuesta producir S/ 840, hay 20 "
                 f"unidades desperdiciadas (S/ 40) y no hay ventas perdidas. Utilidad = "
                 f"1,500 − 840 − 40 − 0 = S/ 620.")

    # ------------------------------- MÉTRICAS -------------------------------
    h1("Métricas")
    parrafo(doc, "La métrica principal es la utilidad neta acumulada: la suma de la utilidad de "
                 "todos los días. Es la que decide qué estrategia gana. Las métricas "
                 "secundarias ayudan a entender por qué:")
    vineta(doc, " total de unidades producidas.", "Producción total:")
    vineta(doc, " total de unidades vendidas.", "Unidades vendidas:")
    vineta(doc, " unidades producidas que no se vendieron.", "Desperdicio total:")
    vineta(doc, " unidades demandadas que no se pudieron atender.", "Ventas perdidas:")
    vineta(doc, " ventas totales / demanda total × 100.", "Nivel de servicio:")

    # ------------------------------ RESULTADOS ------------------------------
    h1("Resultados reales")
    parrafo(doc, f"Estos resultados salen de ejecutar la simulación con los parámetros de la "
                 f"sección de Metodología ({P.dias} días, semilla {P.seed}). No fueron escritos "
                 "a mano: el archivo scripts/generar_informe.py ejecuta el simulador y "
                 "escribe los números en este documento.")
    doc.add_picture(str(tmp / "utilidad.png"), width=Cm(13))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo(doc, "Figura 2. Utilidad acumulada por día.", italica=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(str(tmp / "demanda.png"), width=Cm(13))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo(doc, "Figura 3. Demanda real y producción de cada estrategia.", italica=True).alignment = WD_ALIGN_PARAGRAPH.CENTER

    h1("Tabla comparativa")
    filas = []
    for e in ESTRATEGIAS:
        nombre = f"{e} (ganador)" if e == ganador else e
        filas.append([nombre, f"S/ {formato(m.loc[e, 'Utilidad neta acumulada'])}",
                      formato(m.loc[e, "Desperdicio total"]),
                      formato(m.loc[e, "Ventas perdidas"]),
                      f"{m.loc[e, 'Nivel de servicio (%)']:.1f}%",
                      formato(m.loc[e, "Producción total"])])
    tabla(doc, ["Estrategia", "Utilidad neta acumulada", "Desperdicio", "Ventas perdidas",
                "Nivel de servicio", "Producción total"], filas,
          anchos=[3.2, 3.2, 2.4, 2.4, 2.4, 2.4], resaltar_fila=list(ESTRATEGIAS).index(ganador))
    parrafo(doc, f"Tabla 1. Resultados con {P.dias} días y semilla {P.seed}. Desperdicio y "
                 "ventas perdidas en unidades.", italica=True)

    parrafo(doc, "Validación con varias semillas. ", None).runs[0].bold = True
    parrafo(doc, "Para comprobar que el resultado no depende de una sola secuencia de demanda, se "
                 f"repitió el experimento con 30 semillas consecutivas (de {P.seed} a {P.seed + 29}):")
    filas_v = []
    for e in ESTRATEGIAS:
        v = validacion.loc[e]
        filas_v.append([e, f"S/ {formato(v['Utilidad media'])}", f"S/ {formato(v['Desviación'])}",
                        f"S/ {formato(v['Mínima'])}", f"S/ {formato(v['Máxima'])}",
                        f"{v['Corridas ganadas (%)']:.1f}%"])
    tabla(doc, ["Estrategia", "Utilidad media", "Desviación", "Mínima", "Máxima", "Corridas ganadas"],
          filas_v, anchos=[3, 3, 2.6, 2.6, 2.6, 2.7])

    parrafo(doc, "Cambios de parámetros en vivo. ", None).runs[0].bold = True
    parrafo(doc, "Se repitió la simulación cambiando un solo parámetro a la vez (misma semilla):")
    filas_e = []
    for nombre, prod_media, met in experimentos:
        uu = met["Utilidad neta acumulada"]
        filas_e.append([nombre, f"{prod_media:.1f}", f"S/ {formato(uu['Base'])}",
                        f"S/ {formato(uu['Reactivo'])}", f"S/ {formato(uu['Utilidad'])}",
                        estrategia_ganadora(met)])
    tabla(doc, ["Caso", "Producción media del agente de utilidad", "Utilidad Base",
                "Utilidad Reactivo", "Utilidad Agente de utilidad", "Ganador"],
          filas_e, anchos=[4.6, 2.6, 2.1, 2.1, 2.3, 1.9])

    # ----------------------------- INTERPRETACIÓN -----------------------------
    h1("Interpretación")
    parrafo(doc, f"Con la semilla {P.seed}, la estrategia ganadora es el agente basado en "
                 f"utilidad, con S/ {formato(u['Utilidad'])} de utilidad neta acumulada. "
                 f"Supera al modo base por S/ {formato(ventaja_base)} "
                 f"({ventaja_base / u['Base']:.1%}) y al agente reactivo por "
                 f"S/ {formato(ventaja_reactivo)}.")
    parrafo(doc, f"El modo base produce siempre {P.produccion_base} unidades. Tiene el menor "
                 f"desperdicio ({formato(m.loc['Base', 'Desperdicio total'])} unidades) pero el "
                 f"peor nivel de servicio ({m.loc['Base', 'Nivel de servicio (%)']:.1f}%) y las "
                 f"mayores ventas perdidas ({formato(m.loc['Base', 'Ventas perdidas'])} unidades), "
                 "porque no se adapta cuando la demanda sube a mitad del periodo.")
    parrafo(doc, f"El agente reactivo sí se adapta, y logra el mayor nivel de servicio "
                 f"({m.loc['Reactivo', 'Nivel de servicio (%)']:.1f}%). Pero solo mira el día "
                 "anterior: sube la producción cuando el día anterior se vendió casi todo y se "
                 "queda con sobrantes cuando la demanda baja. Por eso tiene el mayor desperdicio "
                 f"({formato(m.loc['Reactivo', 'Desperdicio total'])} unidades) y gasta mucho en "
                 "producir.")
    parrafo(doc, f"El agente de utilidad equilibra ambos riesgos: su desperdicio "
                 f"({formato(m.loc['Utilidad', 'Desperdicio total'])} unidades) y sus ventas "
                 f"perdidas ({formato(m.loc['Utilidad', 'Ventas perdidas'])} unidades) son "
                 "intermedios. No intenta vender todo ni evitar todo el desperdicio, sino "
                 "maximizar la utilidad, y por eso gana en la métrica principal.")
    parrafo(doc, f"En la validación con 30 semillas, el agente de utilidad ganó "
                 f"{gan_val['Utilidad']:.1f}% de las corridas, el reactivo {gan_val['Reactivo']:.1f}% "
                 f"y el modo base {gan_val['Base']:.1f}%. Esto indica que el resultado no es "
                 "casualidad de una sola semilla.")
    parrafo(doc, f"Sobre los cambios en vivo: al subir el costo de desperdicio de S/ 2 a S/ 8, la "
                 f"producción media del agente de utilidad bajó de {prod0:.1f} a {prod_desp:.1f} "
                 f"unidades. Al subir la penalización por venta perdida de S/ 5 a S/ 15, subió a "
                 f"{prod_pen:.1f}. Es la conducta esperada: más conservador cuando sobrar es caro y "
                 "más generoso cuando faltar es caro.")
    for nombre, _, met in experimentos[1:]:
        if estrategia_ganadora(met) != "Utilidad":
            uu = met["Utilidad neta acumulada"]
            parrafo(doc, f"Ojo: en el caso “{nombre}”, la estrategia ganadora fue "
                         f"{estrategia_ganadora(met)} (S/ {formato(uu.max())}) y no el agente de "
                         f"utilidad (S/ {formato(uu['Utilidad'])}). Cuando perder una venta es muy "
                         "caro conviene producir bastante más que el promedio reciente, y el "
                         "agente de utilidad, que solo mira los últimos 7 días, se queda "
                         "corto frente al reactivo, que sigue subiendo la producción. Ningún "
                         "método es el mejor en todos los escenarios.")
    parrafo(doc, "Limitación importante. ", None).runs[0].bold = True
    parrafo(doc, f"La ventaja del agente de utilidad depende en parte del cambio de nivel de la "
                 f"demanda ({CAMBIO_NIVEL:.0%} a mitad del periodo) que incluimos en la "
                 "simulación. Si la demanda fuera estable alrededor de la producción base, el modo "
                 f"base sería difícil de superar: al repetir las 30 semillas sin ese cambio de "
                 f"nivel, el modo base ganó {gan_val_sc['Base']:.1f}% de las corridas y el agente "
                 f"de utilidad {gan_val_sc['Utilidad']:.1f}%. Es decir, un agente adaptativo aporta "
                 "más valor cuando la demanda cambia. Los datos son sintéticos, no ventas reales.")

    # ----------------------------- CONCLUSIONES -----------------------------
    h1("Conclusiones")
    for texto in [
        f"Se construyó una aplicación funcional que compara tres estrategias de producción sobre "
        "la misma demanda simulada, con KPIs, tabla comparativa y gráficos.",
        f"Con los parámetros base, el agente basado en utilidad obtuvo la mayor utilidad neta "
        f"acumulada (S/ {formato(u['Utilidad'])}), por encima del modo base "
        f"(S/ {formato(u['Base'])}) y del agente reactivo (S/ {formato(u['Reactivo'])}).",
        "Ganar en utilidad no significa ganar en todo: el reactivo tuvo mejor nivel de servicio "
        "y el modo base tuvo menos desperdicio. La utilidad permite equilibrar ambos efectos.",
        "El agente de utilidad responde de forma coherente a los costos: produce menos si el "
        "desperdicio es caro y más si perder una venta es caro.",
        "La ventaja de los agentes depende de que la demanda cambie; con demanda estable el "
        "modo base es una referencia muy competitiva.",
        "La simulación es reproducible: la misma semilla da exactamente el mismo resultado.",
    ]:
        vineta(doc, texto)

    # ---------------------------- TRABAJO FUTURO ----------------------------
    h1("Trabajo futuro")
    parrafo(doc, "Estas ideas no están implementadas en el Avance 1; quedan para la Parte 2 o "
                 "versiones posteriores:")
    vineta(doc, " incorporar Machine Learning para predecir la demanda (por ejemplo regresión "
                "lineal, KNN y árboles de decisión) y alimentar con esa predicción al agente "
                "de utilidad.", "Predicción de demanda:")
    vineta(doc, " usar datos históricos reales (día, producto, promociones, clima o eventos) "
                "en lugar de datos sintéticos.", "Datos reales:")
    vineta(doc, " medir la calidad de la predicción (por ejemplo con MAE) y su impacto "
                "económico.", "Evaluación:")
    vineta(doc, " cargar ventas desde un archivo y entregar recomendaciones diarias de "
                "producción para varios productos.", "Uso en un negocio real:")

    # ------------------------------- USO DE IA -------------------------------
    h1("Uso de IA")
    parrafo(doc, "Claude Code fue utilizado como asistente para generar, integrar, revisar y "
                 "depurar parte del código del proyecto. El equipo revisó, probó y comprende la "
                 "lógica implementada.")

    # --------------------------------- ROLES ---------------------------------
    h1("Roles del equipo")
    parrafo(doc, "El plan inicial del equipo propuso un módulo principal por integrante. Todos "
                 "deben poder explicar el proyecto completo. El aporte real de cada integrante "
                 "(commits en GitHub) está pendiente de confirmar y no se ha completado aquí "
                 "para no inventar información.")
    tabla(doc, ["Integrante", "Módulo propuesto en el plan", "Aporte real / commits"], [
        ["Santiago Valdivia", "Coordinación, integración, métricas, configuración y documentación", "[PENDIENTE]"],
        ["Fabian Cristobal", "Interfaz Streamlit, gráficos y despliegue", "[PENDIENTE]"],
        ["Jesus Camargo", "Simulador y generación de demanda", "[PENDIENTE]"],
        ["Said Taravay", "Lógica de los agentes", "[PENDIENTE]"],
    ], anchos=[3.8, 8, 3.7])

    doc.save(SALIDA)
    return m, validacion, experimentos, val_sin_cambio


if __name__ == "__main__":
    m, validacion, experimentos, val_sin_cambio = construir()
    print(f"Informe generado: {SALIDA}\n")
    print("Resultados (para el README):")
    print(m.round(1).to_string())
    print("\nValidación 30 semillas:")
    print(validacion.round(1).to_string())
    print("\nCambios en vivo (producción media del agente de utilidad):")
    for nombre, prod, met in experimentos:
        print(f"  {nombre}: {prod:.1f}  | utilidades:",
              met["Utilidad neta acumulada"].round(0).to_dict())
    print("\nValidación sin cambio de nivel:")
    print(val_sin_cambio.round(1).to_string())
